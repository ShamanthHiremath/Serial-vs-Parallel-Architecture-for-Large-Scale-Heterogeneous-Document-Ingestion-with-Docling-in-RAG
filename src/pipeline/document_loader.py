"""Document loader module for various file formats."""

import os
from pathlib import Path
from typing import List, Dict, Any
from docx import Document as DocxDocument
import PyPDF2
from pptx import Presentation
from PIL import Image
import logging

logger = logging.getLogger(__name__)


class DocumentLoader:
    """Load and extract text from various document formats."""
    
    SUPPORTED_FORMATS = {'.pdf', '.docx', '.pptx', '.jpg', '.jpeg', '.png', '.bmp', '.tiff'}
    
    def __init__(self):
        """Initialize the document loader."""
        pass
    
    def load_document(self, file_path: str) -> Dict[str, Any]:
        """
        Load a document and extract its content.
        
        Args:
            file_path: Path to the document file
            
        Returns:
            Dictionary containing document metadata and content
        """
        file_path = Path(file_path)
        
        if not file_path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")
        
        extension = file_path.suffix.lower()
        
        if extension not in self.SUPPORTED_FORMATS:
            raise ValueError(f"Unsupported file format: {extension}")
        
        doc_info = {
            'file_path': str(file_path),
            'file_name': file_path.name,
            'file_type': extension,
            'file_size': file_path.stat().st_size,
            'text': '',
            'requires_ocr': False,
            'image_path': None
        }
        
        try:
            if extension == '.pdf':
                doc_info = self._load_pdf(file_path, doc_info)
            elif extension == '.docx':
                doc_info = self._load_docx(file_path, doc_info)
            elif extension == '.pptx':
                doc_info = self._load_pptx(file_path, doc_info)
            elif extension in {'.jpg', '.jpeg', '.png', '.bmp', '.tiff'}:
                doc_info = self._load_image(file_path, doc_info)
        except Exception as e:
            logger.error(f"Error loading document {file_path}: {e}")
            doc_info['error'] = str(e)
        
        return doc_info
    
    def _load_pdf(self, file_path: Path, doc_info: Dict) -> Dict:
        """Load PDF document."""
        try:
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                text_parts = []
                
                for page in pdf_reader.pages:
                    text = page.extract_text()
                    if text:
                        text_parts.append(text)
                
                doc_info['text'] = '\n'.join(text_parts)
                doc_info['page_count'] = len(pdf_reader.pages)
                
                # If no text extracted, mark for OCR
                if not doc_info['text'].strip():
                    doc_info['requires_ocr'] = True
                    doc_info['image_path'] = str(file_path)
        except Exception as e:
            logger.warning(f"PDF text extraction failed for {file_path}, marking for OCR: {e}")
            doc_info['requires_ocr'] = True
            doc_info['image_path'] = str(file_path)
        
        return doc_info
    
    def _load_docx(self, file_path: Path, doc_info: Dict) -> Dict:
        """Load DOCX document."""
        document = DocxDocument(file_path)
        text_parts = [para.text for para in document.paragraphs if para.text.strip()]
        doc_info['text'] = '\n'.join(text_parts)
        doc_info['paragraph_count'] = len(document.paragraphs)
        return doc_info
    
    def _load_pptx(self, file_path: Path, doc_info: Dict) -> Dict:
        """Load PPTX presentation."""
        presentation = Presentation(file_path)
        text_parts = []
        
        for slide in presentation.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    text_parts.append(shape.text)
        
        doc_info['text'] = '\n'.join(text_parts)
        doc_info['slide_count'] = len(presentation.slides)
        return doc_info
    
    def _load_image(self, file_path: Path, doc_info: Dict) -> Dict:
        """Load image file (requires OCR)."""
        # Verify it's a valid image
        try:
            img = Image.open(file_path)
            doc_info['image_size'] = img.size
            doc_info['image_mode'] = img.mode
            img.close()
        except Exception as e:
            raise ValueError(f"Invalid image file: {e}")
        
        doc_info['requires_ocr'] = True
        doc_info['image_path'] = str(file_path)
        return doc_info
    
    def find_documents(self, corpus_path: str) -> List[str]:
        """
        Find all supported documents in a directory.
        
        Args:
            corpus_path: Path to directory containing documents
            
        Returns:
            List of file paths
        """
        corpus_path = Path(corpus_path)
        
        if not corpus_path.exists():
            raise FileNotFoundError(f"Corpus path not found: {corpus_path}")
        
        if corpus_path.is_file():
            return [str(corpus_path)]
        
        documents = []
        for ext in self.SUPPORTED_FORMATS:
            documents.extend(corpus_path.rglob(f'*{ext}'))
        
        return [str(doc) for doc in sorted(documents)]
